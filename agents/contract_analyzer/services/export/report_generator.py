from docx import Document

from agents.contract_analyzer.services.config import (
    SHOW_EXECUTIVE_SUMMARY,
    SHOW_CLAUSE_SUMMARY_COLUMN,
    SHOW_STATUS_COLUMN
)


def export_contract_report(
    executive_summary,
    compliance_results,
    risk_results,
    mitigation_results,
    output_path
):

    doc = Document()

    doc.add_heading(
        "Enterprise Contract Analysis Report",
        level=1
    )

    # ==================================
    # EXECUTIVE SUMMARY
    # ==================================

    if (
        SHOW_EXECUTIVE_SUMMARY
        and executive_summary
    ):

        doc.add_heading(
            "Executive Summary",
            level=2
        )

        doc.add_paragraph(
            f"Overall Risk Rating: "
            f"{executive_summary.get('overall_rating', '')}"
        )

        doc.add_paragraph(
            executive_summary.get(
                "executive_narrative",
                ""
            )
        )

        doc.add_heading(
            "Key Findings",
            level=3
        )

        for finding in executive_summary.get(
            "top_findings",
            []
        ):

            doc.add_paragraph(
                finding,
                style="List Bullet"
            )

        doc.add_heading(
            "Recommendation",
            level=3
        )

        doc.add_paragraph(
            executive_summary.get(
                "recommendation",
                ""
            )
        )

        doc.add_page_break()

    # ==================================
    # TABLE 1
    # ==================================

    doc.add_heading(
        "Table 1 - Compliance Analysis",
        level=2
    )

    cols = 2

    if SHOW_CLAUSE_SUMMARY_COLUMN:
        cols += 1

    if SHOW_STATUS_COLUMN:
        cols += 1

    cols += 1  # Evidence

    table = doc.add_table(
        rows=1,
        cols=cols
    )

    table.style = "Table Grid"

    hdr = table.rows[0].cells

    idx = 0

    hdr[idx].text = "Clause"
    idx += 1

    hdr[idx].text = "Requirement"
    idx += 1

    if SHOW_CLAUSE_SUMMARY_COLUMN:

        hdr[idx].text = (
            "Summary of Key Details"
        )

        idx += 1

    if SHOW_STATUS_COLUMN:

        hdr[idx].text = (
            "Status"
        )

        idx += 1

    hdr[idx].text = (
        "Evidence"
    )

    for row in compliance_results:

        cells = table.add_row().cells

        idx = 0

        cells[idx].text = str(
            row.get(
                "clause",
                ""
            )
        )
        idx += 1

        cells[idx].text = str(
            row.get(
                "requirement",
                ""
            )
        )
        idx += 1

        if SHOW_CLAUSE_SUMMARY_COLUMN:

            cells[idx].text = str(
                row.get(
                    "clause_summary",
                    ""
                )
            )

            idx += 1

        if SHOW_STATUS_COLUMN:

            cells[idx].text = str(
                row.get(
                    "status",
                    ""
                )
            )

            idx += 1

        cells[idx].text = str(
            row.get(
                "evidence",
                ""
            )
        )
    # ==================================
    # TABLE 2
    # ==================================

    doc.add_page_break()

    doc.add_heading(
        "Table 2 - Risk Assessment",
        level=2
    )

    table = doc.add_table(
        rows=1,
        cols=6
    )

    table.style = "Table Grid"

    hdr = table.rows[0].cells

    hdr[0].text = "Clause"
    hdr[1].text = "Requirement"
    hdr[2].text = "RAG"
    hdr[3].text = "Risk"
    hdr[4].text = "Rationale"
    hdr[5].text = "Mitigation"

    for row in risk_results:

        cells = table.add_row().cells

        cells[0].text = str(
            row.get(
                "clause",
                ""
            )
        )

        cells[1].text = str(
            row.get(
                "requirement",
                ""
            )
        )

        cells[2].text = str(
            row.get(
                "rag",
                ""
            )
        )

        cells[3].text = str(
            row.get(
                "risk",
                ""
            )
        )

        cells[4].text = str(
            row.get(
                "rationale",
                ""
            )
        )

        cells[5].text = str(
            row.get(
                "mitigation",
                ""
            )
        )

    # ==================================
    # TABLE 3
    # ==================================

    doc.add_page_break()

    doc.add_heading(
        "Table 3 - Mitigation Checklist",
        level=2
    )

    table = doc.add_table(
        rows=1,
        cols=3
    )

    table.style = "Table Grid"

    hdr = table.rows[0].cells

    hdr[0].text = "Clause"
    hdr[1].text = "Risk"
    hdr[2].text = "Mitigation Recommendation"

    for row in mitigation_results:

        cells = table.add_row().cells

        cells[0].text = str(
            row.get(
                "clause",
                ""
            )
        )

        cells[1].text = str(
            row.get(
                "risk",
                ""
            )
        )

        cells[2].text = str(
            row.get(
                "mitigation",
                ""
            )
        )

    doc.save(
        output_path
    )
    return output_path